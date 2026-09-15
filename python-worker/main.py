"""Optional document parser with PDF/DOCX image extraction."""

from __future__ import annotations

import io
import os
import uuid
import zipfile
from pathlib import Path
from typing import Any
from urllib.parse import quote

import fitz
import httpx
from docx import Document
from pptx import Presentation
from fastapi import FastAPI, File, HTTPException, UploadFile
from markitdown import MarkItDown
from pydantic import BaseModel, Field

ALLOWED_EXT = {".pdf", ".docx", ".pptx"}
MAX_BYTES = 25 * 1024 * 1024
MIN_TEXT_CHARS = 40

app = FastAPI(title="Cogni Document Parser", version="2.0.0")
md = MarkItDown()


class ImageMetadata(BaseModel):
    url: str
    contextText: str = ""
    page: int | None = None
    slide: int | None = None


class ParseResponse(BaseModel):
    fileName: str
    markdown: str
    images: list[ImageMetadata] = Field(default_factory=list)


def storage_config() -> tuple[str, str, str] | None:
    url = os.getenv("SUPABASE_URL", "").rstrip("/")
    key = os.getenv("SUPABASE_SERVICE_ROLE_KEY", "")
    bucket = os.getenv("SUPABASE_STORAGE_BUCKET", "document-images")
    return (url, key, bucket) if url and key else None


def upload_image(
    data: bytes,
    extension: str,
    context: str,
    page: int | None = None,
    slide: int | None = None,
) -> ImageMetadata | None:
    config = storage_config()
    if not config:
        return None

    supabase_url, service_key, bucket = config
    path = f"documents/{uuid.uuid4().hex}{extension}"
    endpoint = f"{supabase_url}/storage/v1/object/{quote(bucket)}/{quote(path)}"
    headers = {
        "Authorization": f"Bearer {service_key}",
        "apikey": service_key,
        "Content-Type": "application/octet-stream",
        "x-upsert": "false",
    }
    try:
        response = httpx.put(endpoint, content=data, headers=headers, timeout=20)
        response.raise_for_status()
    except Exception as error:
        print(f"[parser] image upload failed: {error}")
        return None

    public_url = f"{supabase_url}/storage/v1/object/public/{quote(bucket)}/{quote(path)}"
    return ImageMetadata(url=public_url, contextText=context[:500], page=page, slide=slide)


def extract_pdf(data: bytes) -> tuple[str, list[tuple[bytes, str, str, int | None, int | None]]]:
    document = fitz.open(stream=data, filetype="pdf")
    text_parts: list[str] = []
    images: list[tuple[bytes, str, str, int | None, int | None]] = []
    for page_index, page in enumerate(document, start=1):
        page_text = page.get_text("text")
        text_parts.append(page_text)
        for image in page.get_images(full=True):
            extracted = document.extract_image(image[0])
            if extracted.get("image"):
                extension = f".{extracted.get('ext', 'bin')}"
                images.append((extracted["image"], extension, page_text, page_index, None))
    document.close()
    return "\n\n".join(text_parts), images


def extract_docx(data: bytes) -> tuple[str, list[tuple[bytes, str, str, int | None, int | None]]]:
    document = Document(io.BytesIO(data))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    images: list[tuple[bytes, str, str, int | None, int | None]] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in archive.namelist():
            if not name.startswith("word/media/"):
                continue
            extension = Path(name).suffix.lower() or ".bin"
            images.append((archive.read(name), extension, text, None, None))
    return text, images


def extract_pptx(data: bytes) -> tuple[str, list[tuple[bytes, str, str, int | None, int | None]]]:
    presentation = Presentation(io.BytesIO(data))
    text_parts: list[str] = []
    images: list[tuple[bytes, str, str, int | None, int | None]] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_text = "\n".join(
            shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text
        )
        text_parts.append(f"Slide {slide_number}\n{slide_text}")
        for shape in slide.shapes:
            if not hasattr(shape, "image"):
                continue
            image = shape.image
            images.append((image.blob, f".{image.ext}", slide_text, None, slide_number))
    return "\n\n".join(text_parts), images


def extract_images(items: list[tuple[bytes, str, str, int | None, int | None]]) -> list[ImageMetadata]:
    images: list[ImageMetadata] = []
    for data, extension, context, page, slide in items:
        uploaded = upload_image(data, extension, context, page, slide)
        if uploaded:
            images.append(uploaded)
    return images


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/parse", response_model=ParseResponse)
async def parse(file: UploadFile = File(...)) -> ParseResponse:
    name = file.filename or "upload.bin"
    extension = Path(name).suffix.lower()
    if extension not in ALLOWED_EXT:
        raise HTTPException(status_code=400, detail="Unsupported file type")

    data = await file.read()
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="File too large (max 25MB)")

    try:
        if extension == ".pdf":
            raw_text, raw_images = extract_pdf(data)
        elif extension == ".docx":
            raw_text, raw_images = extract_docx(data)
        else:
            raw_text, raw_images = extract_pptx(data)

        text = raw_text.strip()
        if len(text) < MIN_TEXT_CHARS:
            converted = md.convert(io.BytesIO(data))
            text = (getattr(converted, "text_content", None) or text).strip()
        if len(text) < MIN_TEXT_CHARS:
            raise HTTPException(status_code=422, detail="Could not extract usable text")

        return ParseResponse(
            fileName=name,
            markdown=text,
            images=extract_images(raw_images),
        )
    except HTTPException:
        raise
    except Exception as error:
        print(f"[parser] failed: {error}")
        raise HTTPException(status_code=500, detail="Document parsing failed") from error
